"""
Spreadsheet & CSV Export Service for Batch Pipeline Runs.
Generates styled Excel (.xlsx) workbooks via openpyxl and standard RFC 4180 CSV streams.
"""

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class SpreadsheetReportService:
    """Generates Excel (.xlsx) workbooks and CSV files for pipeline runs."""

    COLUMNS = [
        ("school_name", "School Name", 30),
        ("verification_status", "Verification Status", 20),
        ("overall_score", "DQA Score (%)", 15),
        ("website", "Official Website", 32),
        ("phone", "Telephone", 18),
        ("email", "Email", 26),
        ("address", "Campus Address", 36),
        ("grades", "Grades Offered", 28),
        ("open_enrollment", "Open Enrollment", 16),
        ("extracurricular_programs", "Programs Offered", 30),
        ("description", "Factual Description", 45),
        ("pages_crawled_count", "Pages Crawled", 14),
        ("cache_hits_count", "Cache Hits", 12),
        ("crawled_at", "Executed At", 22),
    ]

    @classmethod
    def generate_excel_workbook(
        cls,
        summary: dict[str, Any],
        schools: list[dict[str, Any]],
    ) -> io.BytesIO:
        """
        Creates a styled Excel workbook from batch run output.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Run Results"

        # Styles
        navy_fill = PatternFill(start_color="0A1F3C", end_color="0A1F3C", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        verified_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
        verified_font = Font(name="Calibri", size=10, bold=True, color="059669")

        review_fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
        review_font = Font(name="Calibri", size=10, bold=True, color="D97706")

        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Title Block
        ws.merge_cells("A1:N1")
        title_cell = ws["A1"]
        title_cell.value = f"SchoolPath AI-DOP • Batch Run Report: {summary.get('batch_name', 'Batch Run')} ({summary.get('job_id')})"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="0A1F3C")
        title_cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30

        # Subtitle / Summary
        ws.merge_cells("A2:N2")
        sub_cell = ws["A2"]
        sub_cell.value = (
            f"Total Schools: {summary.get('total_schools', len(schools))} | "
            f"Verified: {summary.get('verified_count', 0)} | "
            f"Needs Review: {summary.get('needs_review_count', 0)} | "
            f"Average Score: {summary.get('avg_score', 0)}% | "
            f"Date: {summary.get('created_at', '')}"
        )
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="64748B")
        sub_cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 20

        # Empty row 3
        ws.row_dimensions[3].height = 10

        # Header Row (Row 4)
        for col_num, (_, header_title, _) in enumerate(cls.COLUMNS, start=1):
            cell = ws.cell(row=4, column=col_num, value=header_title)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws.row_dimensions[4].height = 28

        # Data Rows (Row 5+)
        for row_idx, s in enumerate(schools, start=5):
            ext = s.get("extracted_fields", {}) or {}
            status = s.get("verification_status", "needs_review")
            is_verified = status == "verified"

            row_data = [
                s.get("school_name", ""),
                "Verified" if is_verified else "Needs Review",
                f"{s.get('overall_score', 0)}%",
                ext.get("website") or s.get("website") or "",
                ext.get("phone") or "",
                ext.get("email") or "",
                ext.get("address") or s.get("address") or "",
                ", ".join(ext.get("grades", []))
                if isinstance(ext.get("grades"), list)
                else str(ext.get("grades") or ""),
                "Yes"
                if ext.get("open_enrollment") is True
                else "No"
                if ext.get("open_enrollment") is False
                else "",
                ", ".join(ext.get("extracurricular_programs", []))
                if isinstance(ext.get("extracurricular_programs"), list)
                else str(ext.get("extracurricular_programs") or ""),
                ext.get("description") or "",
                s.get("pages_crawled_count", 0),
                s.get("cache_hits_count", 0),
                s.get("crawled_at", ""),
            ]

            for col_num, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_num, value=val)
                cell.border = thin_border
                cell.font = regular_font
                cell.alignment = Alignment(vertical="center")

                # Column 2: Status formatting
                if col_num == 2:
                    cell.fill = verified_fill if is_verified else review_fill
                    cell.font = verified_font if is_verified else review_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Column 3: Score formatting
                if col_num == 3:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = bold_font

            ws.row_dimensions[row_idx].height = 22

        # Set Column Widths
        for col_num, (_, _, col_width) in enumerate(cls.COLUMNS, start=1):
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = col_width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def generate_csv_stream(
        cls,
        schools: list[dict[str, Any]],
    ) -> io.StringIO:
        """
        Creates a clean CSV StringIO stream.
        """
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

        # Header Row
        writer.writerow([header for _, header, _ in cls.COLUMNS])

        # Data Rows
        for s in schools:
            ext = s.get("extracted_fields", {}) or {}
            status = s.get("verification_status", "needs_review")
            is_verified = status == "verified"

            writer.writerow(
                [
                    s.get("school_name", ""),
                    "Verified" if is_verified else "Needs Review",
                    s.get("overall_score", 0),
                    ext.get("website") or s.get("website") or "",
                    ext.get("phone") or "",
                    ext.get("email") or "",
                    ext.get("address") or s.get("address") or "",
                    ", ".join(ext.get("grades", []))
                    if isinstance(ext.get("grades"), list)
                    else str(ext.get("grades") or ""),
                    "Yes"
                    if ext.get("open_enrollment") is True
                    else "No"
                    if ext.get("open_enrollment") is False
                    else "",
                    ", ".join(ext.get("extracurricular_programs", []))
                    if isinstance(ext.get("extracurricular_programs"), list)
                    else str(ext.get("extracurricular_programs") or ""),
                    ext.get("description") or "",
                    s.get("pages_crawled_count", 0),
                    s.get("cache_hits_count", 0),
                    s.get("crawled_at", ""),
                ]
            )

        output.seek(0)
        return output
